#!/usr/bin/env python3
"""
fetch_clean.py — Korak 1 izrade NER dataseta za srpsku legislativu.

Preuzima/uvozi tekst zakona, čisti ga, deli na rečenice i snima u JSONL
sa PROVENIJENCIJOM (izvor svake rečenice). Usput ispisuje statistiku i
generiše tabelu izvora (Markdown) za poglavlje "Metodologija" u radu.

Dva režima unosa:
  1) wikisource  — preuzima stranice sa sr.wikisource.org preko API-ja
  2) local       — uvozi lokalne .txt/.html fajlove (npr. skinute sa
                   pravno-informacioni-sistem.rs), autoritativan izvor

Primeri:
  # nađi tačne naslove na Wikisource-u
  python3 fetch_clean.py --search "Закон о раду"

  # preuzmi zakone navedene u laws.json
  python3 fetch_clean.py --mode wikisource --titles-file laws.json \
      --out data/out/legal_raw.jsonl

  # obradi lokalne fajlove (svaki fajl = jedan propis)
  python3 fetch_clean.py --mode local --input-dir data/raw_legal \
      --out data/out/legal_raw.jsonl

Zavisnost: samo `requests`.
"""
import argparse
import csv
import html
import json
import os
import re
import sys
import time
from datetime import date
from urllib.parse import quote

try:
    import requests
except ImportError:
    sys.exit("Nedostaje 'requests'. Instaliraj: pip install requests")

WIKISOURCE_API = "https://sr.wikisource.org/w/api.php"
USER_AGENT = "NER-SR-Legislativa/1.0 (istrazivacki rad; kontakt: student)"

# ---------------------------------------------------------------------------
# Transliteracija ćirilica -> latinica (radi konzistentnosti sa SETimes.SR,
# koji je na latinici). Digrafi (Љ, Њ, Џ) se obrađuju prvi.
# ---------------------------------------------------------------------------
_CYR2LAT_DIGRAPHS = {
    "Љ": "Lj", "љ": "lj", "Њ": "Nj", "њ": "nj", "Џ": "Dž", "џ": "dž",
}
_CYR2LAT = {
    "А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Ђ": "Đ", "Е": "E",
    "Ж": "Ž", "З": "Z", "И": "I", "Ј": "J", "К": "K", "Л": "L", "М": "M",
    "Н": "N", "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T", "Ћ": "Ć",
    "У": "U", "Ф": "F", "Х": "H", "Ц": "C", "Ч": "Č", "Ш": "Š",
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "ђ": "đ", "е": "e",
    "ж": "ž", "з": "z", "и": "i", "ј": "j", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "ћ": "ć",
    "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "č", "ш": "š",
}


def to_latin(text):
    for cyr, lat in _CYR2LAT_DIGRAPHS.items():
        text = text.replace(cyr, lat)
    return "".join(_CYR2LAT.get(ch, ch) for ch in text)


# ---------------------------------------------------------------------------
# Čišćenje teksta
# ---------------------------------------------------------------------------
# Redovi koji su strukturni markeri, a ne rečenice (naslovi članova, glave...)
_STRUCT_LINE = re.compile(
    r"^\s*(члан|члан[а-я]*|glava|глава|одељак|deo|део|"
    r"member|section)\b.*$",
    re.IGNORECASE,
)
# Inline oznaka "Члан 5." / "Član 12." na početku pasusa
_ARTICLE_INLINE = re.compile(r"^\s*(Члан|Član|Чл\.)\s*\d+[а-яa-z]*\.?\s*", re.IGNORECASE)
# Fusnote/ref ostaci, višestruki razmaci
_REF_MARK = re.compile(r"\[\s*\d+\s*\]|\{\{[^}]*\}\}|\[\[[^\]]*\]\]")
_WS = re.compile(r"[ \t ]+")
_MULTINL = re.compile(r"\n{2,}")


def strip_html(text):
    text = re.sub(r"(?is)<script.*?</script>", " ", text)
    text = re.sub(r"(?is)<style.*?</style>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return html.unescape(text)


def read_pdf(path):
    try:
        import pypdf
    except ImportError:
        sys.exit("Za PDF fajlove treba pypdf: pip install pypdf")
    reader = pypdf.PdfReader(path)
    pages = [(p.extract_text() or "") for p in reader.pages]
    text = "\n".join(pages)
    if len(text.strip()) < 50:
        print(f"  ⚠️  {os.path.basename(path)}: skoro nema teksta — možda je "
              f"skeniran PDF (slika). Treba OCR ili kopiraj tekst ručno.",
              file=sys.stderr)
    return text


def read_docx(path):
    try:
        import docx  # python-docx
    except ImportError:
        sys.exit("Za DOCX fajlove treba python-docx: pip install python-docx")
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    # pokupi i tekst iz tabela (zakoni ponekad imaju tabele)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def clean_text(text):
    text = _REF_MARK.sub(" ", text)
    # Podeli na pasuse (prazan red), pa REFLOW: linije unutar pasusa spoji u
    # jedan tok — tako rečenica prelomljena na više vizuelnih redova (čest
    # artefakt PDF/DOCX ekstrakcije) ostaje cela.
    paragraphs = re.split(r"\n\s*\n", text)
    out = []
    for para in paragraphs:
        kept = []
        for ln in para.splitlines():
            ln = _ARTICLE_INLINE.sub("", ln)
            if _STRUCT_LINE.match(ln):
                continue
            ln = ln.strip()
            if ln:
                kept.append(ln)
        if kept:
            out.append(_WS.sub(" ", " ".join(kept)).strip())
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Segmentacija na rečenice (heuristika sa čuvanjem pravnih skraćenica)
# ---------------------------------------------------------------------------
_ABBREV = {
    "чл", "ст", "тач", "бр", "год", "нпр", "итд", "др", "г", "т", "ал",
    "гђа", "проф", "sl", "br", "cl", "st", "tac", "god", "npr", "itd",
    "dr", "tzv", "op", "ur",
}
_SENT_BOUND = re.compile(r"(?<=[.!?])\s+(?=[\"'(]?[A-ZŠĐŽČĆА-Я])")


def split_sentences(text):
    sents = []
    for para in text.split("\n"):
        para = para.strip()
        if not para:
            continue
        candidates = _SENT_BOUND.split(para)
        buf = ""
        for c in candidates:
            piece = (buf + " " + c).strip() if buf else c
            last_word = re.split(r"\s+", piece)[-1].rstrip(".").lower()
            # ako se završava skraćenicom ili golim brojem (npr. "1."),
            # spoji sa sledećim komadom
            if last_word in _ABBREV or re.fullmatch(r"\d+", last_word):
                buf = piece
                continue
            buf = ""
            if len(piece) >= 3:
                sents.append(piece)
        if buf and len(buf) >= 3:
            sents.append(buf)
    return sents


# ---------------------------------------------------------------------------
# Wikisource API
# ---------------------------------------------------------------------------
def api_get(params):
    params = dict(params, format="json")
    r = requests.get(
        WIKISOURCE_API, params=params,
        headers={"User-Agent": USER_AGENT}, timeout=30,
    )
    r.raise_for_status()
    return r.json()


def search_titles(term, limit=25):
    data = api_get({
        "action": "query", "list": "search",
        "srsearch": term, "srlimit": limit,
    })
    return [r["title"] for r in data.get("query", {}).get("search", [])]


def fetch_plaintext(title):
    """Vrati (plaintext, page_url) ili (None, None) ako stranica ne postoji."""
    data = api_get({
        "action": "query", "prop": "extracts|info",
        "explaintext": 1, "exsectionformat": "plain",
        "inprop": "url", "titles": title, "redirects": 1,
    })
    pages = data.get("query", {}).get("pages", {})
    for pid, page in pages.items():
        if pid == "-1" or "missing" in page:
            return None, None
        return page.get("extract", ""), page.get("fullurl", "")
    return None, None


# ---------------------------------------------------------------------------
# Sastavljanje zapisa
# ---------------------------------------------------------------------------
def make_records(sentences, source_meta, translit, start_idx):
    records = []
    for i, sent in enumerate(sentences):
        if translit == "latin":
            sent = to_latin(sent)
        records.append({
            "id": f"{source_meta['slug']}-s{start_idx + i:05d}",
            "text": sent,
            "ner_tags": None,          # popunjava se u fazi anotacije
            "source": source_meta,
        })
    return records


def process_wikisource(titles_file, translit):
    with open(titles_file, encoding="utf-8") as f:
        laws = json.load(f)
    all_records, stats = [], []
    idx = 0
    for law in laws:
        title = law["title"]
        text, url = fetch_plaintext(title)
        if text is None:
            print(f"  ⚠️  NEDOSTAJE na Wikisource: {title}", file=sys.stderr)
            stats.append((title, law.get("slug", ""), 0, "NEDOSTAJE"))
            continue
        cleaned = clean_text(strip_html(text))
        sents = split_sentences(cleaned)
        meta = {
            "type": "zakon",
            "title": title,
            "sl_glasnik": law.get("sl_glasnik", ""),
            "url": url,
            "accessed": date.today().isoformat(),
        }
        recs = make_records(sents, {**meta, "slug": law["slug"]}, translit, idx)
        idx += len(recs)
        all_records.extend(recs)
        stats.append((title, law["slug"], len(recs), "OK"))
        print(f"  ✓ {title}: {len(recs)} rečenica")
        time.sleep(0.5)  # ljubaznost prema API-ju
    return all_records, stats


def process_local(input_dir, translit):
    all_records, stats = [], []
    idx = 0
    files = sorted(
        f for f in os.listdir(input_dir)
        if f.lower().endswith((".txt", ".html", ".htm", ".pdf", ".docx"))
    )
    for fname in files:
        path = os.path.join(input_dir, fname)
        low = fname.lower()
        if low.endswith(".pdf"):
            raw = read_pdf(path)
        elif low.endswith(".docx"):
            raw = read_docx(path)
        else:
            with open(path, encoding="utf-8", errors="replace") as f:
                raw = f.read()
            if low.endswith((".html", ".htm")):
                raw = strip_html(raw)
        cleaned = clean_text(raw)
        sents = split_sentences(cleaned)
        slug = re.sub(r"[^a-z0-9]+", "-", os.path.splitext(fname)[0].lower()).strip("-")
        meta = {
            "type": "zakon",
            "title": os.path.splitext(fname)[0],
            "sl_glasnik": "",   # dopuniti ručno u izlaznom fajlu ako treba
            "url": f"local:{fname}",
            "accessed": date.today().isoformat(),
        }
        recs = make_records(sents, {**meta, "slug": slug}, translit, idx)
        idx += len(recs)
        all_records.extend(recs)
        stats.append((meta["title"], slug, len(recs), "OK"))
        print(f"  ✓ {fname}: {len(recs)} rečenica")
    return all_records, stats


# ---------------------------------------------------------------------------
def write_outputs(records, stats, out_path):
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        for rec in records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    # Tabela izvora (Markdown) — direktno za "Metodologiju"
    md_path = os.path.splitext(out_path)[0] + "_izvori.md"
    total = sum(n for _, _, n, _ in stats)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("| Propis | Oznaka | Rečenica | Status |\n")
        f.write("|---|---|---:|---|\n")
        for title, slug, n, status in stats:
            f.write(f"| {title} | {slug} | {n} | {status} |\n")
        f.write(f"| **UKUPNO** | | **{total}** | |\n")

    print(f"\n=== GOTOVO ===")
    print(f"  Rečenica ukupno: {total}")
    print(f"  JSONL:          {out_path}")
    print(f"  Tabela izvora:  {md_path}")
    if total < 1000:
        print(f"  ⓘ  Za 'dominantno legislativu' cilj je ~1200–1500+ rečenica "
              f"— dodaj još propisa.")


def main():
    ap = argparse.ArgumentParser(description="Preuzmi/očisti/segmentiraj srpsku legislativu za NER.")
    ap.add_argument("--search", metavar="POJAM", help="Pretraži naslove na Wikisource-u i izađi.")
    ap.add_argument("--mode", choices=["wikisource", "local"], help="Izvor unosa.")
    ap.add_argument("--titles-file", default="laws.json", help="JSON lista zakona (wikisource režim).")
    ap.add_argument("--input-dir", default="data/raw_legal", help="Fascikla sa .txt/.html (local režim).")
    ap.add_argument("--out", default="data/out/legal_raw.jsonl", help="Izlazni JSONL.")
    ap.add_argument("--translit", choices=["latin", "none"], default="latin",
                    help="Transliteracija ćirilica->latinica (default: latin, radi sklada sa SETimes.SR).")
    args = ap.parse_args()

    if args.search:
        print(f"Rezultati za '{args.search}':")
        for t in search_titles(args.search):
            print("  -", t)
        return

    if not args.mode:
        ap.error("Zadaj --mode wikisource|local (ili koristi --search).")

    print(f"Režim: {args.mode} | transliteracija: {args.translit}\n")
    if args.mode == "wikisource":
        records, stats = process_wikisource(args.titles_file, args.translit)
    else:
        records, stats = process_local(args.input_dir, args.translit)

    write_outputs(records, stats, args.out)


if __name__ == "__main__":
    main()
